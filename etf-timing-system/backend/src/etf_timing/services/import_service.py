from __future__ import annotations

import csv
import io
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

from etf_timing.config import Settings
from etf_timing.data.cache import SQLiteCache

HISTORY_ALIASES = {
    "日期": "日期",
    "date": "日期",
    "trade_date": "日期",
    "开盘": "开盘",
    "open": "开盘",
    "收盘": "收盘",
    "close": "收盘",
    "最高": "最高",
    "high": "最高",
    "最低": "最低",
    "low": "最低",
    "成交量": "成交量",
    "volume": "成交量",
    "vol": "成交量",
    "成交额": "成交额",
    "amount": "成交额",
    "涨跌幅": "涨跌幅",
    "change_pct": "涨跌幅",
    "换手率": "换手率",
    "turnover_rate": "换手率",
}

SPOT_ALIASES = {
    "代码": "代码",
    "symbol": "代码",
    "code": "代码",
    "名称": "名称",
    "name": "名称",
    "最新价": "最新价",
    "price": "最新价",
    "涨跌幅": "涨跌幅",
    "change_pct": "涨跌幅",
    "成交额": "成交额",
    "amount": "成交额",
    "换手率": "换手率",
    "turnover_rate": "换手率",
    "最新份额": "最新份额",
    "shares": "最新份额",
    "iopv": "IOPV实时估值",
    "iopv实时估值": "IOPV实时估值",
    "基金折价率": "基金折价率",
    "discount_rate": "基金折价率",
    "数据日期": "数据日期",
    "date": "数据日期",
}


def _header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "")


def _scalar(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


class ImportService:
    def __init__(self, cache: SQLiteCache, settings: Settings) -> None:
        self.cache = cache
        self.settings = settings

    def _read_rows(self, filename: str, content: bytes) -> list[dict[str, Any]]:
        suffix = Path(filename).suffix.lower()
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = next(values, None)
            if not headers:
                return []
            names = [str(value or "").strip() for value in headers]
            return [
                dict(zip(names, row, strict=False))
                for row in values
                if any(value is not None for value in row)
            ]
        if suffix == ".csv":
            decoded = None
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if decoded is None:
                raise ValueError("CSV 编码无法识别，请保存为 UTF-8 或 GB18030")
            return list(csv.DictReader(io.StringIO(decoded)))
        raise ValueError("仅支持 .xlsx 或 .csv 文件")

    @staticmethod
    def _normalize(raw_rows: list[dict[str, Any]], aliases: dict[str, str]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for raw in raw_rows:
            normalized: dict[str, Any] = {}
            for key, value in raw.items():
                canonical = aliases.get(_header(key))
                if canonical:
                    normalized[canonical] = _scalar(value)
            if normalized:
                rows.append(normalized)
        return rows

    def import_file(
        self, filename: str, content: bytes, dataset: str, symbol: str
    ) -> dict[str, Any]:
        if len(content) > self.settings.upload_max_bytes:
            raise ValueError(f"文件超过 {self.settings.upload_max_bytes // 1_000_000}MB 限制")
        if not content:
            raise ValueError("上传文件为空")
        source = f"manual-upload/{Path(filename).name}"
        raw_rows = self._read_rows(filename, content)

        if dataset == "history":
            rows = self._normalize(raw_rows, HISTORY_ALIASES)
            required = {"日期", "收盘", "最高", "最低"}
            rows = [row for row in rows if required.issubset(row) and row.get("日期")]
            if not rows:
                raise ValueError("没有有效历史行；至少需要 日期、收盘、最高、最低")
            for row in rows:
                row["日期"] = str(row["日期"])[:10]
                row.setdefault("开盘", row["收盘"])
            rows = list({str(row["日期"]): row for row in rows}.values())
            rows.sort(key=lambda row: str(row["日期"]))
            key = f"etf:history:{symbol.zfill(6)}"
            self.cache.set_import(key, rows, source)
            self.cache.set(key, rows, self.settings.history_ttl_seconds, source)
        elif dataset == "spot":
            rows = self._normalize(raw_rows, SPOT_ALIASES)
            rows = [row for row in rows if row.get("代码") and row.get("最新价") is not None]
            if not rows:
                raise ValueError("没有有效行情行；至少需要 代码、最新价")
            for row in rows:
                row["代码"] = str(row["代码"]).split(".")[0].zfill(6)
                row.setdefault("名称", row["代码"])
                row.setdefault("数据日期", date.today().isoformat())
                row["更新时间"] = datetime.now().astimezone().isoformat()
            key = "etf:spot"
            existing = self.cache.get(key)
            existing_rows = (
                existing.payload if existing and isinstance(existing.payload, list) else []
            )
            merged = {str(row.get("代码") or "").zfill(6): row for row in existing_rows}
            merged.update({str(row["代码"]): row for row in rows})
            rows = list(merged.values())
            self.cache.set_import(key, rows, source)
            self.cache.set(key, rows, self.settings.etf_spot_ttl_seconds, source)
            self.cache.record_etf_spot(rows, source)
        else:
            raise ValueError("dataset 必须是 history 或 spot")

        return {
            "status": "imported",
            "dataset": dataset,
            "symbol": symbol.zfill(6),
            "rows": len(rows),
            "source": source,
        }

    def template(self, dataset: str) -> bytes:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        if dataset == "history":
            sheet.title = "ETF历史行情"
            sheet.append(
                ["日期", "开盘", "收盘", "最高", "最低", "成交量", "成交额", "涨跌幅", "换手率"]
            )
            sheet.append(
                ["2026-08-24", 4.120, 4.156, 4.170, 4.105, 123456789, 512345678, 0.86, 1.25]
            )
        elif dataset == "spot":
            sheet.title = "ETF实时行情"
            sheet.append(
                [
                    "代码",
                    "名称",
                    "最新价",
                    "涨跌幅",
                    "成交额",
                    "换手率",
                    "最新份额",
                    "IOPV实时估值",
                    "基金折价率",
                    "数据日期",
                ]
            )
            sheet.append(
                [
                    "510300",
                    "沪深300ETF",
                    4.156,
                    0.86,
                    512345678,
                    1.25,
                    32000000000,
                    4.158,
                    -0.05,
                    "2026-08-24",
                ]
            )
        else:
            raise ValueError("dataset 必须是 history 或 spot")
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
